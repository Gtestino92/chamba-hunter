from datetime import UTC, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from chamba_hunter.db.connection import Database
from chamba_hunter.db.converters import datetime_to_db
from chamba_hunter.db.migrations import migrate
from chamba_hunter.domain.enums import ApplicationStatus
from chamba_hunter.repositories.application_repository import (
    ApplicationRepository,
)
from chamba_hunter.repositories.job_shortlist_report_repository import (
    JobShortlistReportRepository,
)
from chamba_hunter.services.application_tracking_service import (
    ApplicationTrackingService,
)
from chamba_hunter.services.job_shortlist_report_service import (
    build_summary,
)
from chamba_hunter.services.openoffice_shortlist_actions import (
    add_openoffice_application_actions,
)


NOW = datetime(
    2026,
    9,
    15,
    12,
    0,
    tzinfo=UTC,
)


def _db(value: datetime = NOW) -> str:
    return datetime_to_db(
        value
    )


def _setup_database(
    tmp_path: Path,
) -> Database:
    database = Database(
        tmp_path / "test.db"
    )
    migrate(database)

    with database.transaction() as connection:
        connection.execute(
            """
            INSERT INTO runs (
                id,
                command,
                started_at,
                finished_at,
                status,
                created_by
            )
            VALUES (
                1,
                'prioritize_jobs',
                ?,
                ?,
                'SUCCESS',
                'MANUAL'
            )
            """,
            (
                _db(),
                _db(),
            ),
        )
        connection.execute(
            """
            INSERT INTO search_profiles (
                id,
                name,
                description,
                rules_json,
                is_active,
                created_at,
                updated_at
            )
            VALUES (
                1,
                'BACKEND_SOFTWARE_V1',
                'test',
                '{}',
                1,
                ?,
                ?
            )
            """,
            (
                _db(),
                _db(),
            ),
        )

    return database


def _add_company(
    database: Database,
    company_id: int,
    name: str,
) -> None:
    with database.transaction() as connection:
        connection.execute(
            """
            INSERT INTO companies (
                id,
                name,
                normalized_name,
                status,
                created_at,
                updated_at
            )
            VALUES (?, ?, ?, 'ACTIVE', ?, ?)
            """,
            (
                company_id,
                name,
                name.casefold(),
                _db(),
                _db(),
            ),
        )
        connection.execute(
            """
            INSERT INTO company_ats (
                id,
                company_id,
                provider,
                external_identifier,
                board_url,
                detected_at
            )
            VALUES (?, ?, 'GREENHOUSE', ?, ?, ?)
            """,
            (
                company_id,
                company_id,
                f"company-{company_id}",
                f"https://example.com/{company_id}",
                _db(),
            ),
        )


def _add_ats_job(
    database: Database,
    *,
    job_id: int,
    company_id: int,
    title: str,
) -> None:
    with database.transaction() as connection:
        connection.execute(
            """
            INSERT INTO jobs (
                id,
                company_id,
                company_ats_id,
                external_id,
                title,
                job_url,
                apply_url,
                first_seen_at,
                last_seen_at,
                is_active
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 1)
            """,
            (
                job_id,
                company_id,
                company_id,
                f"ats-{job_id}",
                title,
                f"https://jobs.example/{job_id}",
                f"https://apply.example/{job_id}",
                _db(),
                _db(),
            ),
        )


def _add_lead(
    database: Database,
    *,
    lead_id: int,
    company_id: int,
    title: str,
    canonical_job_id: int | None = None,
) -> None:
    with database.transaction() as connection:
        connection.execute(
            """
            INSERT INTO job_leads (
                id,
                company_id,
                source_type,
                external_id,
                canonical_job_id,
                title,
                job_url,
                apply_url,
                first_seen_at,
                last_seen_at,
                is_active
            )
            VALUES (
                ?,
                ?,
                'GETONBOARD',
                ?,
                ?,
                ?,
                ?,
                ?,
                ?,
                ?,
                1
            )
            """,
            (
                lead_id,
                company_id,
                f"lead-{lead_id}",
                canonical_job_id,
                title,
                f"https://lead.example/{lead_id}",
                f"https://lead-apply.example/{lead_id}",
                _db(),
                _db(),
            ),
        )


def _add_priority(
    database: Database,
    *,
    record_kind: str,
    record_id: int,
    company_id: int,
    company_name: str,
    title: str,
    state: str = "NEW",
    level: str = "VERY_HIGH",
) -> None:
    with database.transaction() as connection:
        connection.execute(
            """
            INSERT INTO job_professional_matches (
                record_kind,
                record_id,
                search_profile_id,
                score,
                match_level,
                role_score,
                skills_score,
                seniority_score,
                leadership_score,
                technology_penalty,
                score_ceiling,
                reasons_json,
                rule_version,
                matched_at
            )
            VALUES (
                ?,
                ?,
                1,
                90,
                ?,
                40,
                25,
                15,
                10,
                0,
                100,
                '{}',
                'MATCHING_TEST',
                ?
            )
            """,
            (
                record_kind,
                record_id,
                level,
                _db(),
            ),
        )
        connection.execute(
            """
            INSERT INTO job_operational_priorities (
                record_kind,
                record_id,
                search_profile_id,
                company_id,
                company_name,
                source_type,
                origin,
                title,
                operational_state,
                professional_score,
                professional_match_level,
                professional_rule_version,
                application_channel,
                application_target,
                job_url,
                apply_url,
                first_seen_at,
                last_seen_at,
                reasons_json,
                rule_version,
                evaluated_at,
                evaluated_run_id
            )
            VALUES (
                ?,
                ?,
                1,
                ?,
                ?,
                ?,
                ?,
                ?,
                ?,
                90,
                ?,
                'MATCHING_TEST',
                'DIRECT_APPLY_URL',
                ?,
                ?,
                ?,
                ?,
                ?,
                '{}',
                'PRIORITY_TEST',
                ?,
                1
            )
            """,
            (
                record_kind,
                record_id,
                company_id,
                company_name,
                (
                    "ATS"
                    if record_kind == "ATS"
                    else "GETONBOARD"
                ),
                (
                    "ATS"
                    if record_kind == "ATS"
                    else "GETONBOARD"
                ),
                title,
                state,
                level,
                f"https://apply.example/{record_id}",
                f"https://jobs.example/{record_id}",
                f"https://apply.example/{record_id}",
                _db(),
                _db(),
                _db(),
            ),
        )


def _add_application(
    database: Database,
    *,
    app_id: int,
    company_id: int,
    record_kind: str,
    record_id: int,
    status: str,
    applied: bool,
) -> None:
    applied_at = (
        _db()
        if applied
        else None
    )
    job_id = (
        record_id
        if record_kind == "ATS"
        else None
    )

    with database.transaction() as connection:
        connection.execute(
            """
            INSERT INTO applications (
                id,
                company_id,
                job_id,
                application_type,
                status,
                applied_at,
                last_status_at,
                created_at,
                updated_at,
                record_kind,
                record_id
            )
            VALUES (
                ?,
                ?,
                ?,
                'JOB',
                ?,
                ?,
                ?,
                ?,
                ?,
                ?,
                ?
            )
            """,
            (
                app_id,
                company_id,
                job_id,
                status,
                applied_at,
                _db(),
                _db(),
                _db(),
                record_kind,
                record_id,
            ),
        )


def _summary(
    database: Database,
):
    return build_summary(
        JobShortlistReportRepository(
            database
        ).load(
            "BACKEND_SOFTWARE_V1"
        )
    )


def _item_by_identity(
    summary,
    record_kind: str,
    record_id: int,
):
    return next(
        item
        for item in summary.source.rows
        if item.record_kind == record_kind
        and item.record_id == record_id
    )


def test_direct_ats_applied_is_history_not_attention(
    tmp_path: Path,
) -> None:
    database = _setup_database(
        tmp_path
    )
    _add_company(database, 1, "DirectCo")
    _add_ats_job(
        database,
        job_id=101,
        company_id=1,
        title="Backend Engineer",
    )
    _add_priority(
        database,
        record_kind="ATS",
        record_id=101,
        company_id=1,
        company_name="DirectCo",
        title="Backend Engineer",
    )
    _add_application(
        database,
        app_id=1,
        company_id=1,
        record_kind="ATS",
        record_id=101,
        status="APPLIED",
        applied=True,
    )

    summary = _summary(database)

    assert not summary.focus
    assert not summary.high_value
    assert not summary.all_current
    assert len(summary.history) == 1
    assert summary.history[0].source.application_status == "APPLIED"
    assert (
        summary.history[0].source.application_applied_at
        is not None
    )


def test_uncanonicalized_lead_applied_is_history_only(
    tmp_path: Path,
) -> None:
    database = _setup_database(
        tmp_path
    )
    _add_company(database, 1, "LeadCo")
    _add_lead(
        database,
        lead_id=201,
        company_id=1,
        title="Backend Engineer",
    )
    _add_priority(
        database,
        record_kind="LEAD",
        record_id=201,
        company_id=1,
        company_name="LeadCo",
        title="Backend Engineer",
    )
    _add_application(
        database,
        app_id=1,
        company_id=1,
        record_kind="LEAD",
        record_id=201,
        status="APPLIED",
        applied=True,
    )

    summary = _summary(database)

    assert not summary.focus
    assert not summary.high_value
    assert not summary.all_current
    assert len(summary.history) == 1


def test_canonicalized_lead_application_suppresses_ats_sections(
    tmp_path: Path,
) -> None:
    database = _setup_database(
        tmp_path
    )
    _add_company(database, 1, "CanonicalCo")
    _add_ats_job(
        database,
        job_id=301,
        company_id=1,
        title="Senior Backend Engineer",
    )
    _add_lead(
        database,
        lead_id=302,
        company_id=1,
        title="Senior Backend Engineer",
        canonical_job_id=301,
    )
    _add_priority(
        database,
        record_kind="ATS",
        record_id=301,
        company_id=1,
        company_name="CanonicalCo",
        title="Senior Backend Engineer",
    )
    _add_application(
        database,
        app_id=1,
        company_id=1,
        record_kind="LEAD",
        record_id=302,
        status="REJECTED",
        applied=True,
    )

    summary = _summary(database)
    row = _item_by_identity(
        summary,
        "ATS",
        301,
    )

    assert row.application_status == "REJECTED"
    assert row.application_applied_at is not None
    assert not summary.focus
    assert not summary.high_value
    assert not summary.all_current
    assert len(summary.history) == 1
    assert summary.history[0].source.record_kind == "ATS"


def test_canonicalized_lead_without_applied_application_does_not_suppress_ats(
    tmp_path: Path,
) -> None:
    database = _setup_database(
        tmp_path
    )
    _add_company(database, 1, "PendingCo")
    _add_ats_job(
        database,
        job_id=401,
        company_id=1,
        title="Backend Engineer",
    )
    _add_lead(
        database,
        lead_id=402,
        company_id=1,
        title="Backend Engineer",
        canonical_job_id=401,
    )
    _add_priority(
        database,
        record_kind="ATS",
        record_id=401,
        company_id=1,
        company_name="PendingCo",
        title="Backend Engineer",
    )

    summary = _summary(database)

    assert len(summary.focus) == 1
    assert len(summary.high_value) == 1
    assert len(summary.all_current) == 1
    assert summary.all_current[0].source.application_applied_at is None


def test_pending_canonical_lead_does_not_suppress_ats(
    tmp_path: Path,
) -> None:
    database = _setup_database(
        tmp_path
    )
    _add_company(database, 1, "PendingLeadCo")
    _add_ats_job(
        database,
        job_id=501,
        company_id=1,
        title="Backend Engineer",
    )
    _add_lead(
        database,
        lead_id=502,
        company_id=1,
        title="Backend Engineer",
        canonical_job_id=501,
    )
    _add_priority(
        database,
        record_kind="ATS",
        record_id=501,
        company_id=1,
        company_name="PendingLeadCo",
        title="Backend Engineer",
    )
    _add_application(
        database,
        app_id=1,
        company_id=1,
        record_kind="LEAD",
        record_id=502,
        status="PENDING",
        applied=False,
    )

    summary = _summary(database)

    assert len(summary.all_current) == 1
    assert summary.all_current[0].source.application_status == "PENDING"
    assert summary.all_current[0].source.application_applied_at is None


def test_any_applied_alias_wins_over_pending_duplicate(
    tmp_path: Path,
) -> None:
    database = _setup_database(
        tmp_path
    )
    _add_company(database, 1, "DuplicateCo")
    _add_ats_job(
        database,
        job_id=601,
        company_id=1,
        title="Backend Engineer",
    )
    _add_lead(
        database,
        lead_id=602,
        company_id=1,
        title="Backend Engineer",
        canonical_job_id=601,
    )
    _add_priority(
        database,
        record_kind="ATS",
        record_id=601,
        company_id=1,
        company_name="DuplicateCo",
        title="Backend Engineer",
    )
    _add_application(
        database,
        app_id=1,
        company_id=1,
        record_kind="LEAD",
        record_id=602,
        status="APPLIED",
        applied=True,
    )
    _add_application(
        database,
        app_id=2,
        company_id=1,
        record_kind="ATS",
        record_id=601,
        status="PENDING",
        applied=False,
    )

    summary = _summary(database)

    assert not summary.all_current
    assert len(summary.history) == 1
    assert summary.history[0].source.application_status == "APPLIED"
    assert (
        summary.history[0].source.application_applied_at
        is not None
    )


def test_no_fuzzy_company_title_application_propagation(
    tmp_path: Path,
) -> None:
    database = _setup_database(
        tmp_path
    )
    _add_company(database, 1, "SameTitleCo")
    _add_ats_job(
        database,
        job_id=701,
        company_id=1,
        title="Backend Engineer",
    )
    _add_ats_job(
        database,
        job_id=702,
        company_id=1,
        title="Backend Engineer",
    )
    _add_lead(
        database,
        lead_id=703,
        company_id=1,
        title="Backend Engineer",
        canonical_job_id=701,
    )
    for job_id in (701, 702):
        _add_priority(
            database,
            record_kind="ATS",
            record_id=job_id,
            company_id=1,
            company_name="SameTitleCo",
            title="Backend Engineer",
        )
    _add_application(
        database,
        app_id=1,
        company_id=1,
        record_kind="LEAD",
        record_id=703,
        status="APPLIED",
        applied=True,
    )

    summary = _summary(database)

    assert [
        item.source.record_id
        for item in summary.all_current
    ] == [
        702,
    ]
    assert [
        item.source.record_id
        for item in summary.history
    ] == [
        701,
    ]


def test_openoffice_lead_tracking_after_canonicalization_is_effective(
    tmp_path: Path,
) -> None:
    database = _setup_database(
        tmp_path
    )
    _add_company(database, 1, "BridgeCo")
    _add_ats_job(
        database,
        job_id=801,
        company_id=1,
        title="Backend Engineer",
    )
    _add_lead(
        database,
        lead_id=802,
        company_id=1,
        title="Backend Engineer",
        canonical_job_id=801,
    )
    _add_priority(
        database,
        record_kind="ATS",
        record_id=801,
        company_id=1,
        company_name="BridgeCo",
        title="Backend Engineer",
    )

    result = ApplicationTrackingService(
        ApplicationRepository(
            database
        )
    ).track_job(
        record_kind="LEAD",
        record_id=802,
        status=ApplicationStatus.APPLIED,
        notes=None,
        notes_provided=False,
    )

    assert result.application.record_kind == "LEAD"

    summary = _summary(database)

    assert not summary.all_current
    assert len(summary.history) == 1
    assert (
        summary.history[0].source.application_applied_at
        is not None
    )


def test_openoffice_actions_keep_record_identity_in_macro_url(
    tmp_path: Path,
) -> None:
    path = tmp_path / "shortlist.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "All Current"
    headers = [
        "Tracked Status",
        "Record Kind",
        "Record ID",
    ]

    for column, header in enumerate(
        headers,
        start=1,
    ):
        sheet.cell(
            row=4,
            column=column,
            value=header,
        )

    sheet.cell(
        row=5,
        column=2,
        value="LEAD",
    )
    sheet.cell(
        row=5,
        column=3,
        value=802,
    )
    workbook.save(path)

    assert (
        add_openoffice_application_actions(
            path
        )
        == 1
    )

    updated = load_workbook(path)
    cell = updated[
        "All Current"
    ].cell(
        row=5,
        column=1,
    )

    assert "kind=LEAD" in cell.hyperlink.target
    assert "id=802" in cell.hyperlink.target
